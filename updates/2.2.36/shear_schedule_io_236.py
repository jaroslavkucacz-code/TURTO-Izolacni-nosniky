from __future__ import annotations

"""File intake and auditable row parser for the shear-dowel Decoder.

No capacity calculation takes place here. Raster input always requires review.
PDF text/word positions are preferred; scans are never silently dropped.
"""
import csv
from dataclasses import dataclass
import io
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import unicodedata
from typing import Any, Callable


@dataclass(frozen=True)
class FileText:
    text: str
    source: str
    review_required: bool = False
    notice: str = ""


def _plain(text: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", text.lower()) if not unicodedata.combining(c))


def words_to_text(words: list[tuple]) -> str:
    """Group by physical baseline, not PDF object or OCR reading order."""
    lines: list[list[tuple]] = []
    for word in sorted(words, key=lambda w: ((w[1] + w[3]) / 2, w[0])):
        if not str(word[4]).strip():
            continue
        cy = (word[1] + word[3]) / 2
        line = next((r for r in reversed(lines[-3:])
                     if abs(cy - (r[0][1] + r[0][3]) / 2)
                     <= max(2, min(word[3]-word[1], r[0][3]-r[0][1]) * .48)), None)
        if line is None:
            lines.append([word])
        else:
            line.append(word)
    out = []
    for line in lines:
        line.sort(key=lambda w: w[0])
        text = str(line[0][4])
        for left, right in zip(line, line[1:]):
            separation = right[0] - left[2]
            text += ("\t" if separation > max(12, (left[3]-left[1])*1.6) else " ") + str(right[4])
        out.append(text)
    return "\n".join(out)


def _ocr_image(path: Path) -> str:
    """Local-only OCR; Windows built-in recognizer, optional existing Tesseract."""
    from PIL import Image, ImageOps, ImageDraw
    with tempfile.TemporaryDirectory(prefix="turto_ocr_") as folder:
        target = Path(folder) / "input.png"
        with Image.open(path) as original:
            image = ImageOps.exif_transpose(original).convert("RGB")
            if image.width * image.height > 50_000_000:
                raise ValueError("Obrázek je příliš velký. Vyberte výřez tabulky.")
            # Rules often join all cells into one OCR block. Remove only long,
            # almost-solid table rules in the working copy, never source pixels.
            grey = image.convert("L")
            w,h = image.size
            rows = [y for y in range(h) if sum(grey.crop((0,y,w,y+1)).histogram()[:180]) > w*.60]
            cols = [x for x in range(w) if sum(grey.crop((x,0,x+1,h)).histogram()[:180]) > h*.60]
            draw = ImageDraw.Draw(image)
            for y in rows:
                draw.rectangle((0,max(0,y-1),w,min(h,y+1)),fill="white")
            for x in cols:
                draw.rectangle((max(0,x-1),0,min(w,x+1),h),fill="white")
            scale = min(2.0, 2400 / max(image.size))
            image = image.resize((max(1, round(image.width*scale)), max(1, round(image.height*scale))))
            image.save(target)
        kwargs = dict(capture_output=True, timeout=75, check=True)
        if os.name == "nt":
            script = Path(__file__).with_name("hsd_windows_ocr.ps1")
            env = dict(os.environ, TURTO_OCR_INPUT=str(target))
            result = subprocess.run(["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(script)],
                                    env=env, creationflags=0x08000000, **kwargs)
            data = json.loads(result.stdout.decode("utf-8-sig").strip())
            words = [(w["x"], w["y"], w["x"]+w["w"], w["y"]+w["h"], w["text"]) for w in data]
        elif shutil.which("tesseract"):
            result = subprocess.run(["tesseract", str(target), "stdout", "-l", "eng", "--psm", "6", "tsv"], **kwargs)
            data = csv.DictReader(io.StringIO(result.stdout.decode("utf-8")), delimiter="\t")
            words = [(int(w["left"]), int(w["top"]), int(w["left"])+int(w["width"]),
                      int(w["top"])+int(w["height"]), w["text"]) for w in data if w["text"].strip()]
        else:
            raise RuntimeError("OCR není dostupné. Použijte textové PDF, Excel nebo text ze schránky.")
        text = words_to_text(words)
        if not text.strip():
            raise ValueError("V obrázku nebyl nalezen čitelný text. Vyberte ostřejší výřez tabulky.")
        return text


def read_file(path: str | Path, *, page: int | None = None) -> FileText:
    p = Path(path).resolve()
    if not p.is_file():
        raise ValueError("Soubor nebyl nalezen.")
    if p.stat().st_size > 80_000_000:
        raise ValueError("Soubor je příliš velký; vyberte samostatný výkaz nebo výřez.")
    suffix = p.suffix.lower()
    if suffix in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}:
        return FileText(_ocr_image(p), str(p), True, "OCR: zkontrolujte označení, písmeno V a počty. Text lze opravit před vložením.")
    if suffix == ".pdf":
        import fitz
        with fitz.open(p) as doc:
            if doc.needs_pass:
                raise ValueError("PDF je chráněné heslem.")
            if page is None and len(doc) > 1:
                raise ValueError("Vyberte číslo stránky s výkazem.")
            n = 0 if page is None else page-1
            if n < 0 or n >= len(doc):
                raise ValueError(f"PDF má {len(doc)} stran; zadejte platné číslo stránky.")
            words = doc[n].get_text("words")
            text = words_to_text(words)
            if text.strip():
                return FileText(text, f"{p} | strana {n+1}")
            # Only the explicitly selected page is sent to the local recognizer.
            with tempfile.TemporaryDirectory(prefix="turto_scan_") as folder:
                target = Path(folder) / "page.png"
                doc[n].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(target)
                text = _ocr_image(target)
            return FileText(text, f"{p} | strana {n+1}", True, "Sken PDF / OCR: potvrďte kontrolu označení a počtů.")
    if suffix in {".xlsx", ".xlsm"}:
        # Runtime dependency already used by TURTO; never execute macros.
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            sheet = wb.active
            if sheet.max_row > 20000 or sheet.max_column > 100:
                raise ValueError("Vyberte menší výkaz (nejvýše 20 000 řádků a 100 sloupců).")
            rows = ["\t".join("" if v is None else str(v) for v in row).rstrip() for row in sheet.iter_rows(values_only=True)]
            return FileText("\n".join(rows), f"{p} | list {sheet.title}", False, f"Načten aktivní list: {sheet.title}.")
        finally:
            wb.close()
    if suffix in {".txt", ".csv", ".tsv"}:
        raw = p.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1250")
        if suffix == ".csv":
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
                text = "\n".join("\t".join(row) for row in csv.reader(io.StringIO(text), dialect))
            except csv.Error:
                pass
        return FileText(text, str(p))
    raise ValueError("Podporováno: PDF, XLSX/XLSM, CSV/TXT/TSV a PNG/JPG/BMP/TIF.")


def parse_rows(text: str, *, decoder: Callable, defaults: dict[str, str], existing_names: set[str], use_defaults: bool = False):
    from shear_dowels_schedule import ScheduleItem
    names = set(existing_names)
    header: dict[str, int] = {}
    out = []
    for line, raw in enumerate(text.splitlines(), 1):
        raw = raw.strip()
        if not raw or re.fullmatch(r"[\s|:+\-=]+", raw):
            continue
        cells = [v.strip() for v in re.split(r"\t|\||;", raw.strip("|"))]
        plain = [_plain(c) for c in cells]
        if re.fullmatch(r"vykaz\s+smykovych\s+[^0-9]+", _plain(raw)):
            continue
        product = next((i for i,v in enumerate(plain) if v in {"nazev", "oznaceni", "typ", "vyrobek", "designation", "name"}), None)
        quantity = next((i for i,v in enumerate(plain) if v in {"pocet", "ks", "mnozstvi", "quantity", "qty"}), None)
        if product is not None and quantity is not None:
            header = {"product": product, "quantity": quantity}
            continue
        if re.fullmatch(r"(?:nazev|oznaceni|typ)\s+(?:pocet|ks|mnozstvi)", _plain(raw)):
            continue
        if re.fullmatch(r"(?:celkem|total)\s*[:\t ]*\d+\s*(?:ks)?", _plain(raw)):
            continue
        count = "1"
        designation = raw
        explicit_qty = False
        position = ""
        if header and len(cells) > max(header.values()):
            designation, count = cells[header["product"]], cells[header["quantity"]]
            explicit_qty = True
        elif len(cells) >= 3 and re.fullmatch(r"[A-Za-z][\w./-]*",cells[0]) and not decoder(cells[0]):
            from shear_dowels_schedule import _position_qty_content
            position, qty_old, designation = _position_qty_content(raw, "")
            count, explicit_qty = str(qty_old), True
        elif len(cells) > 1 and decoder(cells[0]):
            designation, count = cells[0], cells[-1]
            explicit_qty = True
        else:
            prefix = re.fullmatch(r"(\d+)\s*[x×]\s*(.+)", raw, re.I)
            trailing = re.fullmatch(r"(.+?)\s+([+-]?\d+(?:[.,]\d+)?)(?:\s*ks)?", raw, re.I)
            if prefix and decoder(prefix[2]):
                count, designation, explicit_qty = prefix[1], prefix[2], True
            elif trailing and decoder(trailing[1]):
                designation, count, explicit_qty = trailing[1], trailing[2], True
        index = 1
        while f"S{index:03}" in names:
            index += 1
        name = position or f"S{index:03}"
        names.add(name)
        valid_count = re.fullmatch(r"\s*([1-9]\d*)\s*(?:ks)?\s*", count, re.I)
        qty = int(valid_count[1]) if valid_count else 0
        info = decoder(designation)
        geometry = dict(defaults) if use_defaults else {"slab":"", "gap":"", "concrete":""}
        for key, pattern in {"slab":r"\bh\s*[:=]?\s*([+-]?\d+(?:[.,]\d+)?)\s*mm", "gap":r"\b(?:sp[aá]ra|gap|joint)\s*[:=]?\s*([+-]?\d+(?:[.,]\d+)?)\s*mm", "concrete":r"\b(C\s*\d+\s*/\s*\d+)\b"}.items():
            match = re.search(pattern, raw, re.I)
            if match:
                geometry[key] = match[1].replace(" ", "").replace(",", ".")
        try:
            h, gap = float(geometry.get("slab") or 0), float(geometry.get("gap") or 0)
        except ValueError:
            h = gap = -1
        known = bool(geometry.get("slab") and geometry.get("gap") and geometry.get("concrete"))
        values = {"designation": designation, "slab_mm":h, "gap_mm":gap, "concrete":geometry.get("concrete", ""),
                  "geometry_confirmed":known, "quantity_explicit":explicit_qty, "source_text":raw}
        item = ScheduleItem(line, raw, name, qty, values)
        if not info:
            item.error = "Označení nebylo rozpoznáno; opravte vstupní text."
        elif not valid_count:
            item.error = "Počet musí být celé kladné číslo."
        elif not math.isfinite(h) or not math.isfinite(gap) or h < 0 or gap < 0 or (known and h == 0):
            item.error = "Neplatná tloušťka nebo šířka spáry."
        else:
            values.update({key: info.get(key, "") for key in ("manufacturer", "family", "size", "movement")})
            values["canonical_designation"] = info.get("designation", designation)
            item.result = str(values["canonical_designation"])
            if not known:
                item.result += " • pouze typ a počet; doplňte h, spáru a beton"
            if not explicit_qty:
                item.result += " • počet neuveden: 1 ks"
        out.append(item)
    return out
